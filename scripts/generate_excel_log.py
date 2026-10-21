import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_ai_log():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "AI Interaction Log"
    
    header_fill = PatternFill(start_color="00FF00", end_color="00FF00", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="000000")
    
    regular_font = Font(name="Arial", size=10, color="000000")
    bold_font = Font(name="Arial", size=10, bold=True, color="000000")
    
    thin_border = Border(
        left=Side(style='thin', color='000000'),
        right=Side(style='thin', color='000000'),
        top=Side(style='thin', color='000000'),
        bottom=Side(style='thin', color='000000')
    )
    
    headers = [
        "№п.п",
        "Компонент системи",
        "Назва задачі",
        "Хто розробляв\n(Самостійно / AI / Спільно)",
        "Коментар / Промт"
    ]
    
    ws.append(headers)
    
    rows = [
        (
            1,
            "ТЗ та аналіз вимог",
            "Опис технічного завдання",
            "Самостійно",
            "— Формулювання функціональних вимог для варіанта 20 (датасети 5, 10, 20 точок), структури 6 графіків, вимог числової стійкості (відмова від нормальних рівнянь) та критеріїв точності."
        ),
        (
            2,
            "Математичний апарат",
            "Теоретичний базис полінома Лагранжа",
            "Самостійно",
            "— Математичний вивід інтерполяційного многочлена Лагранжа L_k(x), обґрунтування вибору глобальної інтерполяції та визначення умов запобігання діленню на нуль у вузлах."
        ),
        (
            3,
            "Математичне ядро",
            "Алгоритм інтерполяції Лагранжа (math.ts)",
            "Самостійно",
            "— Самостійна реалізація алгоритму lagrangeInterpolate з обчислювальною складністю O(n^2) мовою TypeScript; забезпечення точного проходження кривої через вузли g(x_i) = y_i."
        ),
        (
            4,
            "Математичний апарат",
            "Обґрунтування стійкого МНК через QR",
            "Самостійно",
            "— Аналіз проблеми квадратичного зростання числа обумовленості kappa(X^T*X) = kappa(X)^2. Прийняття рішення про використання стійкої Householder QR-факторизації замість схеми Гаусса."
        ),
        (
            5,
            "Математичне ядро",
            "Алгоритм МНК-апроксимації (fitPolynomialOLS)",
            "Самостійно",
            "— Побудова матриці Вандермонда для поліномів ступенів m=2, 3, розв'язання перевизначеної системи через ml-matrix QrDecomposition та розрахунок залишків r_i = y_i - y_fit."
        ),
        (
            6,
            "Математичне ядро",
            "Розрахунок аналітичних метрик якості",
            "Самостійно",
            "— Реалізація функцій розрахунку MSE, RMSE, MAE, коефіцієнта детермінації R^2 та максимального відхилення MaxError для порівняльної оцінки адекватності моделей."
        ),
        (
            7,
            "Архітектура системи",
            "Базовий каркас React-компонентів",
            "Спільно",
            "Промт: 'Створи базовий шаблон розмітки компонентів дашборду: Header, ControlSidebar, MetricsGrid, MainChart'.\nКоментар: AI надав шаблонний каркас JSX-розмітки, студент реалізував збереження стану (useState), типізацію та логіку взаємодії."
        ),
        (
            8,
            "UI / UX та стилі",
            "Стилізація Tailwind CSS (Modern Light Theme)",
            "AI",
            "Промт: 'Стилізуй інтерфейс у світлій темі: чистий білий фон, картки з тонкою рамкою border-slate-200, прибери неоновий glow та ALL CAPS, дизайн під shadcn/ui'.\nКоментар: AI підібрав класи Tailwind CSS, палітру та радіуси заокруглень."
        ),
        (
            9,
            "UI / UX та стилі",
            "Верстка панелі керування (ControlSidebar)",
            "Спільно",
            "Промт: 'Зверстай радіо-кнопки вибору датасету (5, 10, 20 точок) та перемикач відображення кривих'.\nКоментар: AI згенерував розмітку елементів форми, студент додав валідацію вхідних значень та обробники зміни параметрів."
        ),
        (
            10,
            "Графічний рушій",
            "Математика полотна Canvas та проекція координат",
            "Самостійно",
            "— Розробка математики лінійного перетворення світ/екран: алгоритми computeViewport, toScreen, toWorld, розрахунок динамічного кроку осей Nice Numbers, масштабування під Retina devicePixelRatio."
        ),
        (
            11,
            "Графічний рушій",
            "Рендеринг кривих, точок та залишків (MainChart.tsx)",
            "Спільно",
            "Промт: 'Надай синтаксичний шаблон Canvas 2D API для малювання штрихової лінії залишків та маркерів точок'.\nКоментар: AI надав приклад викликів Canvas API, студент реалізував повну математику рендерингу кривих та hit-testing курсору."
        ),
        (
            12,
            "Інтерактивність",
            "Анімація покрокової побудови та морфінгу",
            "Спільно",
            "Промт: 'Напиши шаблон циклу requestAnimationFrame для нарощування прогресу анімації t від 0 до 1'.\nКоментар: AI надав структуру таймінгу, студент запрограмував покрокове додавання вузлів та математичний морфінг полінома від прямої meanY."
        ),
        (
            13,
            "Візуалізація",
            "Компонент матриці 6 графіків (Matrix6View.tsx)",
            "Спільно",
            "Промт: 'Створи розмітку CSS Grid 2x3 з Tailwind для розміщення 6 карточок з міні-графіками'.\nКоментар: AI створив структуру CSS-сітки, студент підключив незалежний розрахунок аналітичних моделей для 5, 10 та 20 точок."
        ),
        (
            14,
            "Тестування та аналітика",
            "Аналітичне тестування та феномен Рунге",
            "Самостійно",
            "— Проведення обчислювальних експериментів на 3 датасетах, формування зведеної таблиці похибок, аналіз осциляцій полінома 19-го ступеня при N=20 та обґрунтування стійкості МНК."
        ),
        (
            15,
            "Документування",
            "Формування аналітичного звіту та висновків",
            "Самостійно",
            "— Повний теоретичний опис, відповіді на 10 контрольних запитань, математичний аналіз отриманих результатів, формулювання інженерних висновків та оформлення звіту."
        )
    ]
    
    for row in rows:
        ws.append(row)
        
    for col_idx in range(1, 6):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
        
    for row_idx in range(2, len(rows) + 2):
        ws.row_dimensions[row_idx].height = 46
        for col_idx in range(1, 6):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            cell.font = regular_font
            if col_idx == 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in (2, 3):
                cell.alignment = Alignment(horizontal="left", vertical="center")
                cell.font = bold_font
            elif col_idx == 4:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                val = str(cell.value)
                if val == "Самостійно":
                    cell.font = Font(name="Arial", size=10, bold=True, color="008000")
                elif val == "AI":
                    cell.font = Font(name="Arial", size=10, bold=True, color="0000FF")
                else:
                    cell.font = Font(name="Arial", size=10, bold=True, color="B8860B")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                
    ws.row_dimensions[1].height = 32
    col_widths = {1: 8, 2: 24, 3: 28, 4: 26, 5: 65}
    for col_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width
        
    try:
        wb.save("reports/AI_Interaction_Log.xlsx")
        print("reports/AI_Interaction_Log.xlsx updated successfully!")
    except PermissionError:
        alt_path = "reports/AI_Interaction_Log_v2.xlsx"
        wb.save(alt_path)
        print(f"reports/AI_Interaction_Log.xlsx is currently open in Excel. Saved updated copy to {alt_path}!")

if __name__ == "__main__":
    generate_ai_log()

